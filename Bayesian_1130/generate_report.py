import os
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import re
from datetime import datetime, timezone

# ----------------------------
# 配置：使用相对路径（假设此脚本位于 Bayesian_1130 目录下）
# ----------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
APRIORI_DIR = os.path.join(BASE_DIR, "result", "apriori_results")
BAYESIAN_DIR = os.path.join(BASE_DIR, "result", "bayesian_results")

# 动态生成报告文件名
today_str = datetime.now().strftime('%Y-%m-%d')
OUTPUT_DOCX = os.path.join(BASE_DIR, f"装备使用质量评价分析报告_{today_str}.docx")


# ----------------------------
# 辅助函数：添加带标题的图片
# ----------------------------
def add_image_with_caption(doc, img_path, caption_text):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = p.add_run()
        try:
            run.add_picture(img_path, width=Inches(6.0))
        except Exception as e:
            p.text = f"[图片加载失败: {e}]"
        
        caption_p = doc.add_paragraph(caption_text, style='Caption')
        caption_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        doc.add_paragraph()  # 空行
    else:
        doc.add_paragraph(f"[图片缺失: {img_path}]", style='Intense Quote')


# ----------------------------
# 解析 prediction_report.txt
# ----------------------------
def parse_prediction_report(report_path):
    if not os.path.exists(report_path):
        return None
    with open(report_path, "r", encoding="utf-8") as f:
        content = f.read()

    data = {}
    # 提取总体准确度
    acc_match = re.search(r"总体准确度:\s*([0-9.]+)", content)
    data["accuracy"] = float(acc_match.group(1)) if acc_match else None

    # 提取每个类别的 precision, recall, f1, support
    class_blocks = re.findall(
        r"([\u4e00-\u9fa5\w\s]+?):\s*?\n\s*支持样本数.*?(\d+\.?\d*).*?\n\s*精确率.*?([0-9.]+).*?\n\s*召回率.*?([0-9.]+).*?\n\s*F1分数.*?([0-9.]+)",
        content, re.DOTALL
    )
    classes = []
    for name, sup, prec, rec, f1 in class_blocks:
        classes.append({
            "name": name.strip(),
            "support": int(float(sup)),
            "precision": float(prec),
            "recall": float(rec),
            "f1-score": float(f1)
        })
    data["classes"] = classes
    return data


# ----------------------------
# 主程序：生成 Word 报告
# ----------------------------
def main():
    doc = Document()
    
    # 使用标准UTC时间设置文件内部元数据，避免时区问题
    utc_now = datetime.now(timezone.utc)
    doc.core_properties.created = utc_now
    doc.core_properties.modified = utc_now

    # 使用本地时间在报告中显示，符合用户习惯
    local_now = datetime.now()

    # 标题
    doc.add_heading('装备使用质量评价分析报告', 0)
    p = doc.add_paragraph(f"报告生成日期: {local_now.strftime('%Y-%m-%d %H:%M:%S')}")
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph()  # 空行

    # ============ Apriori 数据挖掘板块 ============
    doc.add_heading('一、Apriori 数据挖掘板块', level=1)
    doc.add_paragraph("本部分对不同离散化方法在Apriori算法中的性能进行了全面评估，包括生成规则的数量、执行时间以及综合性能对比。")

    apriori_images = [
        ("故障预测规则提升度.png", "图1-1 故障预测规则提升度"),
        ("离散化方法规则数量对比.png", "图1-2 各离散化方法生成规则数量对比"),
        ("离散化方法执行时间对比.png", "图1-3 各离散化方法执行时间对比"),
        ("离散化方法性能综合对比.png", "图1-4 离散化方法性能综合对比 (执行时间 vs 规则数量)"),
    ]

    for filename, caption in apriori_images:
        img_path = os.path.join(APRIORI_DIR, filename)
        add_image_with_caption(doc, img_path, caption)

    # ============ 贝叶斯网络板块 ============
    doc.add_heading('二、贝叶斯网络板块', level=1)
    doc.add_paragraph("本部分展示了基于贝叶斯网络模型的设备状态预测结果，包括网络结构、模型性能评估和详细分类报告。")

    # 插入 BN 结构图
    bn_img = os.path.join(BAYESIAN_DIR, "bn_structure.png")
    add_image_with_caption(doc, bn_img, "图2-1 贝叶斯网络结构图")

    # 插入混淆矩阵
    cm_img = os.path.join(BAYESIAN_DIR, "confusion_matrix.png")
    add_image_with_caption(doc, cm_img, "图2-2 模型预测结果混淆矩阵")

    # 解析并插入预测报告
    report_path = os.path.join(BAYESIAN_DIR, "prediction_report.txt")
    report_data = parse_prediction_report(report_path)

    if report_data:
        doc.add_heading("2.1 模型性能评估报告", level=2)
        doc.add_paragraph(f"模型的整体准确率为 **{report_data['accuracy']:.4f}**。下表为详细分类报告：")

        # 创建表格
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = '类别'
        hdr_cells[1].text = '精确率 (Precision)'
        hdr_cells[2].text = '召回率 (Recall)'
        hdr_cells[3].text = 'F1分数 (F1-Score)'
        hdr_cells[4].text = '支持样本数 (Support)'

        for cell in hdr_cells:
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER


        for cls in report_data["classes"]:
            row_cells = table.add_row().cells
            row_cells[0].text = cls["name"]
            row_cells[1].text = f"{cls['precision']:.2f}"
            row_cells[2].text = f"{cls['recall']:.2f}"
            row_cells[3].text = f"{cls['f1-score']:.2f}"
            row_cells[4].text = str(cls["support"])
            for cell in row_cells:
                cell.paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        doc.add_paragraph()
        doc.add_paragraph("**结论:**")
        doc.add_paragraph("模型在“正常运行”状态上表现出色，精确率达到100%，但在少数类（如“散热系统故障”）上精确率较低，存在一定的误报风险。总体而言，模型具有较高的召回率，能够有效识别出绝大多数的故障状态。")
    else:
        doc.add_paragraph("[报告文件缺失或无法解析]", style='Intense Quote')

    # 保存文档
    doc.save(OUTPUT_DOCX)
    print(f"✅ 报告已成功生成: {OUTPUT_DOCX}")

if __name__ == "__main__":
    main()